from __future__ import annotations

import sys
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))

from converter import OcrLine, OcrWord, PdfToWordConverter  # noqa: E402


class FakeTable:
    cells = [(0.0, 0.0, 60.0, 20.0), (60.0, 0.0, 180.0, 20.0)]


class ConverterTests(unittest.TestCase):
    def test_normalize_chinese_ocr_spacing(self) -> None:
        self.assertEqual(PdfToWordConverter._normalize_text("\u667a \u80fd \u5236 \u9020 \uff0c \u5de5 \u7a0b"), "\u667a\u80fd\u5236\u9020\uff0c\u5de5\u7a0b")
        self.assertEqual(PdfToWordConverter._normalize_text("PDF \u8f6c Word"), "PDF \u8f6c Word")

    def test_cell_indices_selects_source_grid_range(self) -> None:
        edges = [0.0, 60.0, 180.0]
        self.assertEqual(PdfToWordConverter._cell_indices(edges, 60.0, 180.0), (1, 1))
        self.assertEqual(PdfToWordConverter._cell_indices(edges, 0.0, 180.0), (0, 1))

    def test_append_table_writes_editable_cell_text(self) -> None:
        document = Document()
        lines = (
            OcrLine("\u59d3\u540d", 10.0, 5.0, 20.0, 8.0, ()),
            OcrLine("\u66fe\u6587\u745e", 75.0, 5.0, 30.0, 8.0, ()),
        )
        PdfToWordConverter()._append_table(document, FakeTable(), lines, 1.0, 1.0)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(document.tables[0].cell(0, 0).text, "\u59d3\u540d")
        self.assertEqual(document.tables[0].cell(0, 1).text, "\u66fe\u6587\u745e")


    def test_words_split_across_adjacent_cells_by_coordinates(self) -> None:
        line = OcrLine("\u6027\u522b\u7537", 10.0, 5.0, 100.0, 8.0, (
            OcrWord("\u6027", 10.0, 5.0, 12.0, 8.0),
            OcrWord("\u522b", 25.0, 5.0, 12.0, 8.0),
            OcrWord("\u7537", 75.0, 5.0, 12.0, 8.0),
        ))
        converter = PdfToWordConverter()
        self.assertEqual(converter._text_in_box((line,), (0.0, 0.0, 60.0, 20.0), 1.0, 1.0), "\u6027\u522b")
        self.assertEqual(converter._text_in_box((line,), (60.0, 0.0, 120.0, 20.0), 1.0, 1.0), "\u7537")

if __name__ == "__main__":
    unittest.main()
