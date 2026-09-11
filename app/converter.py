from __future__ import annotations

import json
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Iterable

import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt


ProgressCallback = Callable[[str], None]


@dataclass(frozen=True)
class OcrWord:
    text: str
    x: float
    y: float
    width: float
    height: float

    @property
    def center(self) -> tuple[float, float]:
        return self.x + self.width / 2, self.y + self.height / 2


@dataclass(frozen=True)
class OcrLine:
    text: str
    x: float
    y: float
    width: float
    height: float
    words: tuple[OcrWord, ...]

    @property
    def center(self) -> tuple[float, float]:
        return self.x + self.width / 2, self.y + self.height / 2


class ConversionError(RuntimeError):
    """Raised when the input cannot be turned into a trustworthy DOCX."""


class PdfToWordConverter:
    """OCR PDF pages and reconstruct bordered tables as editable Word tables."""

    def __init__(self, language_tag: str = "zh-Hans-CN") -> None:
        self.language_tag = language_tag
        self._ocr_script = Path(__file__).with_name("windows_ocr.ps1")

    def convert(self, source: Path, destination: Path, progress: ProgressCallback) -> None:
        source = source.resolve()
        destination.parent.mkdir(parents=True, exist_ok=True)
        if source.suffix.lower() != ".pdf":
            raise ConversionError("只支持 PDF 文件。")
        if not source.is_file():
            raise ConversionError(f"找不到输入文件：{source}")

        document = self._new_document()
        try:
            with pdfplumber.open(source) as plumber_pdf:
                rendered_pdf = pdfium.PdfDocument(str(source))
                if len(plumber_pdf.pages) != len(rendered_pdf):
                    raise ConversionError("PDF 页数读取不一致，已停止转换以避免遗漏内容。")

                for page_index, page in enumerate(plumber_pdf.pages):
                    progress(f"正在识别第 {page_index + 1}/{len(plumber_pdf.pages)} 页…")
                    image_path, image_size = self._render_page(rendered_pdf, page_index)
                    try:
                        lines = self._ocr(image_path)
                    finally:
                        image_path.unlink(missing_ok=True)

                    self._append_page(document, page, lines, image_size)
                    if page_index < len(plumber_pdf.pages) - 1:
                        document.add_page_break()
        except ConversionError:
            raise
        except Exception as exc:
            raise ConversionError(f"转换失败：{exc}") from exc

        document.save(str(destination))
        progress(f"已生成：{destination.name}")

    @staticmethod
    def _new_document() -> Document:
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(1.7)
        section.bottom_margin = Cm(1.7)
        section.left_margin = Cm(1.7)
        section.right_margin = Cm(1.7)
        normal = document.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        normal.font.size = Pt(10.5)
        return document

    @staticmethod
    def _render_page(pdf: pdfium.PdfDocument, page_index: int) -> tuple[Path, tuple[int, int]]:
        page = pdf[page_index]
        bitmap = page.render(scale=2.4, rotation=0)
        image = bitmap.to_pil()
        handle = tempfile.NamedTemporaryFile(prefix="pdf-to-word-", suffix=".png", delete=False)
        path = Path(handle.name)
        handle.close()
        image.save(path)
        return path, image.size

    def _ocr(self, image_path: Path) -> tuple[OcrLine, ...]:
        powershell = shutil.which("powershell.exe") or shutil.which("powershell")
        if not powershell:
            raise ConversionError("此软件需要 Windows PowerShell 才能调用系统 OCR。")
        command = [
            powershell,
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(self._ocr_script),
            "-ImagePath",
            str(image_path),
            "-LanguageTag",
            self.language_tag,
        ]
        result = subprocess.run(command, text=True, capture_output=True, encoding="utf-8", errors="replace")
        if result.returncode:
            message = result.stderr.strip() or result.stdout.strip() or "Windows OCR 没有返回结果。"
            raise ConversionError(message)
        try:
            payload = json.loads(result.stdout)
        except json.JSONDecodeError as exc:
            raise ConversionError(f"无法读取 Windows OCR 结果：{result.stdout[:200]}") from exc
        return tuple(
            OcrLine(
                text=self._normalize_text(item["text"]),
                x=float(item["x"]),
                y=float(item["y"]),
                width=float(item["width"]),
                height=float(item["height"]),
                words=tuple(
                    OcrWord(
                        text=self._normalize_text(word["text"]),
                        x=float(word["x"]),
                        y=float(word["y"]),
                        width=float(word["width"]),
                        height=float(word["height"]),
                    )
                    for word in item.get("words", [])
                    if word["text"].strip()
                ),
            )
            for item in payload.get("lines", [])
            if item["text"].strip()
        )

    @staticmethod
    def _normalize_text(text: str) -> str:
        """Remove OCR-inserted gaps between Chinese characters without joining English words."""
        compact = " ".join(text.strip().split())
        cjk_or_punctuation = r"[\u3000-\u303f\u3400-\u9fff\uff00-\uffef]"
        return re.sub(rf"(?<={cjk_or_punctuation})\s+(?={cjk_or_punctuation})", "", compact)

    def _append_page(self, document: Document, page, lines: tuple[OcrLine, ...], image_size: tuple[int, int]) -> None:
        tables = page.find_tables()
        scale_x = image_size[0] / page.width
        scale_y = image_size[1] / page.height
        table_rectangles = [table.bbox for table in tables]

        outside_lines = [
            line
            for line in lines
            if not any(self._inside(line.center, bbox, scale_x, scale_y) for bbox in table_rectangles)
        ]

        table_events = [(table.bbox[1], "table", table) for table in tables]
        line_events = [(line.y / scale_y, "line", line) for line in outside_lines]
        for _, kind, item in sorted(table_events + line_events, key=lambda event: event[0]):
            if kind == "table":
                self._append_table(document, item, lines, scale_x, scale_y)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.add_run(item.text)

    def _append_table(self, document: Document, table, lines: tuple[OcrLine, ...], scale_x: float, scale_y: float) -> None:
        x_edges = self._edges(cell[0] for cell in table.cells) + self._edges(cell[2] for cell in table.cells)
        y_edges = self._edges(cell[1] for cell in table.cells) + self._edges(cell[3] for cell in table.cells)
        x_edges = sorted(set(x_edges))
        y_edges = sorted(set(y_edges))
        if len(x_edges) < 2 or len(y_edges) < 2:
            return

        word_table = document.add_table(rows=len(y_edges) - 1, cols=len(x_edges) - 1)
        word_table.style = "Table Grid"
        word_table.autofit = False
        for row in word_table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        for source_cell in table.cells:
            x0, y0, x1, y1 = source_cell
            start_row, end_row = self._cell_indices(y_edges, y0, y1)
            start_col, end_col = self._cell_indices(x_edges, x0, x1)
            target = word_table.cell(start_row, start_col)
            if start_row != end_row or start_col != end_col:
                target = target.merge(word_table.cell(end_row, end_col))
            target.width = Cm(max((x1 - x0) * 0.03528, 0.4))
            content = self._text_in_box(lines, (x0, y0, x1, y1), scale_x, scale_y)
            target.text = content
            for paragraph in target.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)

    @staticmethod
    def _edges(values: Iterable[float]) -> list[float]:
        return [round(value, 2) for value in values]

    @staticmethod
    def _cell_indices(edges: list[float], start: float, end: float) -> tuple[int, int]:
        start_index = min(range(len(edges)), key=lambda index: abs(edges[index] - start))
        end_index = min(range(len(edges)), key=lambda index: abs(edges[index] - end)) - 1
        return start_index, max(start_index, end_index)

    @staticmethod
    def _inside(point: tuple[float, float], box: tuple[float, float, float, float], scale_x: float, scale_y: float) -> bool:
        x, y = point[0] / scale_x, point[1] / scale_y
        return box[0] <= x <= box[2] and box[1] <= y <= box[3]

    def _text_in_box(self, lines: tuple[OcrLine, ...], box: tuple[float, float, float, float], scale_x: float, scale_y: float) -> str:
        words = [word for line in lines for word in line.words if self._inside(word.center, box, scale_x, scale_y)]
        if words:
            return self._words_to_text(words)
        selected = [line for line in lines if self._inside(line.center, box, scale_x, scale_y)]
        return "\n".join(line.text for line in sorted(selected, key=lambda line: (line.y, line.x)))

    def _words_to_text(self, words: list[OcrWord]) -> str:
        rows: list[list[OcrWord]] = []
        row_centers: list[float] = []
        for word in sorted(words, key=lambda item: (item.center[1], item.x)):
            center_y = word.center[1]
            if not rows or abs(center_y - row_centers[-1]) > max(4.0, word.height * 0.65):
                rows.append([word])
                row_centers.append(center_y)
                continue
            rows[-1].append(word)
            row_centers[-1] = sum(item.center[1] for item in rows[-1]) / len(rows[-1])
        return "\n".join(
            self._normalize_text(" ".join(word.text for word in sorted(row, key=lambda item: item.x)))
            for row in rows
        )


def qn(name: str) -> str:
    """Resolve an OOXML namespace name without exposing python-docx internals elsewhere."""
    from docx.oxml.ns import qn as qualify_name

    return qualify_name(name)
